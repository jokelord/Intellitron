from __future__ import annotations

import zipfile
from contextlib import contextmanager
from dataclasses import replace

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorksheetImage
from PIL import Image

from backend.config import ReadFileConfig
from backend.tools.read_file.errors import ReadFileError
from backend.tools.read_file.paddleocr_client import OCRPage
from backend.tools.read_file.service import ReadFileService
from backend.tools.read_file.spreadsheet_media_extractor import extract_embedded_images
from backend.tools.read_file.spreadsheet_parser import parse_xlsx


def _workbook(path, image_path=None):
    workbook = Workbook()
    first = workbook.active
    first.title = "中文表"
    first.append(["名称", "值"])
    first.append(["合并", "内容"])
    first.merge_cells("A2:B2")
    first["C3"] = "=1+2"
    second = workbook.create_sheet("Second")
    second.append(["B"])
    if image_path:
        first.add_image(WorksheetImage(image_path), "B4")
        first.add_image(WorksheetImage(image_path), "D8")
    workbook.save(path)


def test_markitdown_xlsx_preserves_sheet_order_names_and_tables(tmp_path):
    path = tmp_path / "book.xlsx"
    _workbook(path)
    sheets = parse_xlsx(path)
    assert [sheet.name for sheet in sheets] == ["中文表", "Second"]
    assert "名称" in sheets[0].markdown
    assert "合并" in sheets[0].markdown
    assert "=1+2" in sheets[0].markdown or "3" in sheets[0].markdown


def test_ooxml_media_extraction_maps_anchors_and_deduplicates_sha(tmp_path):
    bitmap = tmp_path / "embedded.png"
    Image.new("RGB", (16, 16), "red").save(bitmap)
    path = tmp_path / "book.xlsx"
    _workbook(path, bitmap)

    images = extract_embedded_images(path, ReadFileConfig())
    try:
        assert len(images) == 1
        assert images[0].sha256
        assert images[0].placements == (("中文表", "B4"), ("中文表", "D8"))
        assert images[0].path.read_bytes().startswith(b"\x89PNG")
    finally:
        for image in images:
            image.close()


def test_malformed_drawing_relationship_is_rejected(tmp_path):
    path = tmp_path / "bad.xlsx"
    _workbook(path)
    rewritten = tmp_path / "rewritten.xlsx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rewritten, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"worksheets/sheet1.xml", b"../../../escape.xml")
            target.writestr(info, data)
    with pytest.raises(ReadFileError, match="SPREADSHEET_INVALID"):
        extract_embedded_images(rewritten, ReadFileConfig())


class _OCR:
    def __init__(self):
        self.calls = 0

    def parse_file(self, path, file_type, options, check_abort):
        self.calls += 1
        return [OCRPage(1, "embedded OCR text")]


def test_xlsx_service_inserts_each_anchor_but_ocrs_unique_media_once(tmp_path):
    bitmap = tmp_path / "embedded.png"
    Image.new("RGB", (16, 16), "red").save(bitmap)
    path = tmp_path / "book.xlsx"
    _workbook(path, bitmap)
    ocr = _OCR()
    class Guard:
        checks = 0

        @contextmanager
        def acquire(self, check_abort):
            yield

        def check_memory(self):
            self.checks += 1

    guard = Guard()
    service = ReadFileService(
        config=replace(ReadFileConfig(), allowed_roots=[str(tmp_path)]),
        ocr_client=ocr,
        resource_guard=guard,
    )
    output = service.read(str(path))
    assert "### 嵌入图片：中文表!B4" in output
    assert "### 嵌入图片：中文表!D8" in output
    assert output.count("embedded OCR text") == 2
    assert "source_sheet=中文表" in output
    assert ocr.calls == 1
    assert guard.checks >= 1


def test_xlsx_bounds_are_enforced_before_markitdown(tmp_path, monkeypatch):
    path = tmp_path / "book.xlsx"
    _workbook(path)
    called = False

    def forbidden(*args, **kwargs):
        nonlocal called
        called = True
        raise AssertionError("MarkItDown must not run")

    monkeypatch.setattr("markitdown.MarkItDown.convert_local", forbidden)
    with pytest.raises(ReadFileError, match="SPREADSHEET_INVALID"):
        parse_xlsx(
            path,
            replace(
                ReadFileConfig(),
                max_xlsx_sheets=1,
                max_xlsx_cells=100,
                max_xlsx_rows=100,
                max_xlsx_columns=100,
            ),
        )
    assert called is False


def test_xlsx_stream_validates_forged_row_and_cell_coordinates_before_markitdown(
    tmp_path, monkeypatch
):
    source = tmp_path / "oversized.xlsx"
    workbook = Workbook()
    workbook.active["A200001"] = "outside configured row limit"
    workbook.save(source)
    forged = tmp_path / "forged-dimension.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(forged, "w") as target:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="A1:A200001"', b'ref="A1"')
                data = data.replace(b'ref="A200001:A200001"', b'ref="A1"')
                data = data.replace(b'ref="A200001"', b'ref="A1"')
                assert b'<dimension ref="A1"' in data
                assert b'<row r="200001"' in data
                assert b'<c r="A200001"' in data
            target.writestr(info, data)

    called = False

    def forbidden(*args, **kwargs):
        nonlocal called
        called = True
        raise AssertionError("MarkItDown must not inspect an out-of-range XLSX")

    monkeypatch.setattr("markitdown.MarkItDown.convert_local", forbidden)
    with pytest.raises(ReadFileError, match="SPREADSHEET_INVALID"):
        parse_xlsx(forged, ReadFileConfig())
    assert called is False
