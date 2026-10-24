from __future__ import annotations

import re
import zipfile
from pathlib import Path

from defusedxml import ElementTree as SafeElementTree
from openpyxl import load_workbook

from backend.config import ReadFileConfig

from .errors import ReadFileError, ReadFileErrorCode
from .models import SheetResult


def _column_number(value: str) -> int:
    result = 0
    for character in value:
        if character.isalpha():
            result = result * 26 + ord(character.upper()) - 64
    return result


def _cell_coordinate(value: str) -> tuple[int, int]:
    match = re.fullmatch(r"\$?([A-Za-z]+)\$?([0-9]+)", value)
    if match is None:
        raise ValueError("invalid XLSX cell coordinate")
    row = int(match.group(2))
    column = _column_number(match.group(1))
    if row <= 0 or column <= 0:
        raise ValueError("invalid XLSX cell coordinate")
    return row, column


def _validate_coordinate(value: str, config: ReadFileConfig) -> None:
    row, column = _cell_coordinate(value)
    if row > config.max_xlsx_rows or column > config.max_xlsx_columns:
        raise ReadFileError(
            ReadFileErrorCode.SPREADSHEET_INVALID,
            "XLSX worksheet coordinates exceed configured limits",
        )


def _validate_xlsx_bounds(source: Path, config: ReadFileConfig) -> None:
    try:
        with zipfile.ZipFile(source) as archive:
            workbook = SafeElementTree.fromstring(archive.read("xl/workbook.xml"))
            sheet_count = len(
                workbook.findall(
                    ".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"
                )
            )
            if sheet_count > config.max_xlsx_sheets:
                raise ReadFileError(
                    ReadFileErrorCode.SPREADSHEET_INVALID,
                    "XLSX sheet count exceeds configured limit",
                )
            total_cells = 0
            for name in archive.namelist():
                if not name.startswith("xl/worksheets/") or not name.endswith(".xml"):
                    continue
                with archive.open(name) as worksheet_stream:
                    for event, element in SafeElementTree.iterparse(
                        worksheet_stream, events=("start", "end")
                    ):
                        local_name = element.tag.rsplit("}", 1)[-1]
                        if event == "start" and local_name == "dimension":
                            reference = element.attrib.get("ref", "A1")
                            _validate_coordinate(reference.split(":")[-1], config)
                        elif event == "start" and local_name == "row":
                            raw_row = element.attrib.get("r")
                            if raw_row is None or not raw_row.isdigit():
                                raise ReadFileError(
                                    ReadFileErrorCode.SPREADSHEET_INVALID,
                                    "XLSX row coordinate is malformed",
                                )
                            row = int(raw_row)
                            if row <= 0 or row > config.max_xlsx_rows:
                                raise ReadFileError(
                                    ReadFileErrorCode.SPREADSHEET_INVALID,
                                    "XLSX row coordinate exceeds configured limit",
                                )
                        elif event == "start" and local_name == "c":
                            reference = element.attrib.get("r")
                            if reference is None:
                                raise ReadFileError(
                                    ReadFileErrorCode.SPREADSHEET_INVALID,
                                    "XLSX cell coordinate is missing",
                                )
                            _validate_coordinate(reference, config)
                            total_cells += 1
                            if total_cells > config.max_xlsx_cells:
                                raise ReadFileError(
                                    ReadFileErrorCode.SPREADSHEET_INVALID,
                                    "XLSX cell count exceeds configured limit",
                                )
                        if event == "end":
                            element.clear()
    except ReadFileError:
        raise
    except (OSError, KeyError, ValueError, SyntaxError, zipfile.BadZipFile) as exc:
        raise ReadFileError(
            ReadFileErrorCode.SPREADSHEET_INVALID,
            "XLSX structure could not be validated",
        ) from exc


def parse_xlsx(
    path: str | Path, config: ReadFileConfig | None = None
) -> list[SheetResult]:
    source = Path(path)
    config = config or ReadFileConfig()
    _validate_xlsx_bounds(source, config)
    workbook = None
    cached_workbook = None
    try:
        from markitdown import MarkItDown

        converted = MarkItDown(enable_plugins=False).convert_local(
            source, file_extension=".xlsx"
        )
        markdown = converted.text_content
        workbook = load_workbook(source, read_only=True, data_only=False)
        cached_workbook = load_workbook(source, read_only=True, data_only=True)
        sheet_names = list(workbook.sheetnames)
        formulas: dict[str, list[str]] = {}
        for worksheet in workbook.worksheets:
            cached_sheet = cached_workbook[worksheet.title]
            sheet_formulas = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cached = cached_sheet[cell.coordinate].value
                        suffix = f" -> {cached}" if cached is not None else ""
                        sheet_formulas.append(
                            f"- Formula {cell.coordinate}: `{cell.value}`{suffix}"
                        )
            formulas[worksheet.title] = sheet_formulas
    except Exception as exc:
        raise ReadFileError(
            ReadFileErrorCode.SPREADSHEET_INVALID,
            "MarkItDown could not parse the XLSX workbook",
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()
        if cached_workbook is not None:
            cached_workbook.close()
    sections: list[SheetResult] = []
    for index, name in enumerate(sheet_names):
        marker = re.compile(rf"^##\s+{re.escape(name)}\s*$", re.MULTILINE)
        match = marker.search(markdown)
        if match is None:
            content = f"## {name}"
        else:
            next_start = len(markdown)
            for later_name in sheet_names[index + 1 :]:
                later = re.search(
                    rf"^##\s+{re.escape(later_name)}\s*$",
                    markdown[match.end() :],
                    re.MULTILINE,
                )
                if later is not None:
                    next_start = match.end() + later.start()
                    break
            content = markdown[match.start() : next_start].strip()
        if formulas.get(name):
            content += "\n\n### Formulas\n" + "\n".join(formulas[name])
        sections.append(SheetResult(name=name, index=index, markdown=content))
    return sections
